# -*- coding: utf-8 -*-
"""
[EN] Build Documentation/Overall_Validation.xlsx: the one overall validation
     workbook. It does NOT copy tests; it indexes the 13 module workbooks
     (test counts, stage-1 list, status snapshot) with links. Refresh after
     testers change statuses:  python3 tools/make_overall_validation.py
[FA] ساخت فایل ولیدیشن کلی: نمای کلی همهٔ ماژول‌ها با لینک و شمارش وضعیت‌ها.
     تست‌ها اینجا تکرار نمی‌شوند؛ جزئیات و ثبت نتیجه در فایل هر ماژول است.
"""
import pathlib
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / "Documentation" / "Overall_Validation.xlsx"
FONT = "Vazirmatn"
TITLE, HEAD, BAND, IDCOL = "1F3864", "1F4E79", "F2F6FC", "D9E2F3"
INPUT = "FFF2CC"
THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUSES = ["در انتظار", "انجام شد", "قبول", "رد", "مسدود", "نیاز به تکرار"]

MODULES = [
    ("CalLut", "کالیبراسیون (CalLut)", "CalLut_Validation.xlsx"),
    ("Changeover", "قطع‌ووصل خروجی (Changeover)", "Changeover_Board_Validation.xlsx"),
    ("Charger", "شارژر (Charger)", "Charger_Validation.xlsx"),
    ("Diag", "لاگ تشخیصی (Diag)", "Diag_Validation.xlsx"),
    ("EspLink", "ارتباط با پنل (EspLink)", "EspLink_Validation.xlsx"),
    ("Fault", "تشخیص خطا (Fault)", "Fault_Validation.xlsx"),
    ("Imbalance", "عدم‌توازن باتری (Imbalance)", "Imbalance_Validation.xlsx"),
    ("Jitter", "مقایسه‌گر جریان (Jitter)", "Jitter_Validation.xlsx"),
    ("McuPowerPath", "تغذیهٔ میکرو (McuPowerPath)", "McuPowerPath_Board_Validation.xlsx"),
    ("Measurement", "اندازه‌گیری (Measurement)", "Measurement_Board_Validation.xlsx"),
    ("Protection", "حفاظت (Protection)", "Protection_Validation.xlsx"),
    ("Rtos", "سیستم‌عامل و تسک‌ها (Rtos)", "Rtos_Validation.xlsx"),
    ("Ui", "چراغ و بوق (UI)", "UI_Board_Validation.xlsx"),
]
LAYERS = ["برنامه تست", "تست‌های برنامه‌نویسی", "تست‌های آزمایشگاهی"]


def find_header(ws):
    for r in range(1, 8):
        vals = {c.column: (str(c.value).strip() if c.value else "") for c in ws[r]}
        if any("شناسه" in v or v == "ردیف" for v in vals.values()):
            return r, {v: col for col, v in vals.items() if v}
    return None, {}


def sheet_stats(path):
    """counts per layer + priority/status counts of the practical sheet."""
    wb = openpyxl.load_workbook(path, data_only=False)
    counts, top, status = {}, 0, {s: 0 for s in STATUSES}
    for layer in LAYERS:
        n = 0
        if layer in wb.sheetnames:
            ws = wb[layer]
            hdr, heads = find_header(ws)
            if hdr:
                for r in range(hdr + 1, ws.max_row + 1):
                    if ws.cell(r, 1).value is not None:
                        n += 1
                        if layer == "برنامه تست":
                            if heads.get("اولویت") and ws.cell(r, heads["اولویت"]).value == "بالا":
                                top += 1
                            if heads.get("وضعیت"):
                                v = ws.cell(r, heads["وضعیت"]).value
                                if v in status:
                                    status[v] += 1
        counts[layer] = n
    return counts, top, status


def stage1_rows(path, module_fa):
    wb = openpyxl.load_workbook(path)
    out = []
    if "برنامه تست" not in wb.sheetnames:
        return out
    ws = wb["برنامه تست"]
    hdr, heads = find_header(ws)
    if not hdr or "اولویت" not in heads:
        return out
    for r in range(hdr + 1, ws.max_row + 1):
        if ws.cell(r, heads["اولویت"]).value == "بالا":
            out.append((str(ws.cell(r, 1).value), module_fa,
                        str(ws.cell(r, heads.get("دسته", 2)).value or "")))
    return out


def style_title(ws, text, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(1, 1, text)
    c.font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=TITLE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row, col)
        c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def build():
    wb = openpyxl.Workbook()

    # ---------- نمای کلی ----------
    ws = wb.active
    ws.title = "نمای کلی"
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "1F4E79"
    heads = ["ردیف", "ماژول", "فایل ولیدیشن", "عملی", "برنامه‌نویسی", "آزمایشگاهی",
             "«بالا» در عملی"] + STATUSES + ["فایل ماشین حالت"]
    style_title(ws, "نمای کلی ولیدیشن همهٔ ماژول‌ها — جزئیات و ثبت نتیجه در فایل هر ماژول است", len(heads))
    for i, h in enumerate(heads, start=1):
        ws.cell(2, i, h)
    style_header(ws, 2, len(heads))
    r = 3
    totals_from = r
    stage1 = []
    for mod, fa, fname in MODULES:
        path = ROOT / "Firmware" / "Modules" / mod / fname
        counts, top, status = sheet_stats(path)
        stage1 += [(tid, fa, cat, fname) for tid, fa2, cat in
                   [(t[0], t[1], t[2]) for t in stage1_rows(path, fa)]]
        link = f"../Firmware/Modules/{mod}/{fname}"
        sm = f"../Firmware/Modules/{mod}/{mod}_State_Machine.xlsx"
        if mod == "Ui":
            sm = "../Firmware/Modules/Ui/Ui_State_Machine.xlsx"
        row = [r - 2, fa, "باز کردن فایل", counts["برنامه تست"], counts["تست‌های برنامه‌نویسی"],
               counts["تست‌های آزمایشگاهی"], top] + [status[s] for s in STATUSES] + ["ماشین حالت"]
        for i, v in enumerate(row, start=1):
            c = ws.cell(r, i, v)
            c.font = Font(name=FONT, size=11)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER
            c.fill = PatternFill("solid", fgColor=IDCOL if i == 2 else (BAND if r % 2 == 1 else "FFFFFF"))
        lc = ws.cell(r, 3)
        lc.hyperlink = link
        lc.font = Font(name=FONT, size=11, color="0563C1", underline="single")
        sc = ws.cell(r, len(heads))
        sc.hyperlink = sm
        sc.font = Font(name=FONT, size=11, color="0563C1", underline="single")
        r += 1
    # system row
    spath = ROOT / "Documentation" / "System_State_Scenarios.xlsx"
    swb = openpyxl.load_workbook(spath)
    tc = swb["Test_Cases"]
    n_sys = sum(1 for row in range(4, tc.max_row + 1) if tc.cell(row, 1).value)
    sstat = {s: 0 for s in STATUSES}
    for row in range(4, tc.max_row + 1):
        v = tc.cell(row, 8).value
        if v in sstat:
            sstat[v] += 1
    row = [r - 2, "تست‌های سطح سیستم", "باز کردن فایل", n_sys, 0, 0, 0] + \
          [sstat[s] for s in STATUSES] + ["—"]
    for i, v in enumerate(row, start=1):
        c = ws.cell(r, i, v)
        c.font = Font(name=FONT, size=11)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        c.fill = PatternFill("solid", fgColor=IDCOL if i == 2 else (BAND if r % 2 == 1 else "FFFFFF"))
    lc = ws.cell(r, 3)
    lc.hyperlink = "System_State_Scenarios.xlsx"
    lc.font = Font(name=FONT, size=11, color="0563C1", underline="single")
    r += 1
    # totals (live SUM)
    ws.cell(r, 2, "جمع").font = Font(name=FONT, size=11, bold=True)
    for col in list(range(4, 8)) + list(range(8, 14)):
        L = get_column_letter(col)
        c = ws.cell(r, col, f"=SUM({L}{totals_from}:{L}{r - 1})")
        c.font = Font(name=FONT, size=11, bold=True)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    widths = [6, 30, 16, 9, 13, 13, 12] + [11] * 6 + [14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ---------- مرحله اول ----------
    ws2 = wb.create_sheet("مرحله اول")
    ws2.sheet_view.rightToLeft = True
    ws2.sheet_properties.tabColor = "C00000"
    style_title(ws2, "تست‌های مرحلهٔ اول (اولویت «بالا») — اول همین‌ها روی برد انجام می‌شود", 5)
    for i, h in enumerate(["ردیف", "شناسه", "ماژول", "دسته", "فایل ولیدیشن"], start=1):
        ws2.cell(2, i, h)
    style_header(ws2, 2, 5)
    r = 3
    for tid, fa, cat, fname in stage1:
        mod = next(m for m, f2, fn in MODULES if fn == fname)
        vals = [r - 2, tid, fa, cat, "باز کردن فایل"]
        for i, v in enumerate(vals, start=1):
            c = ws2.cell(r, i, v)
            c.font = Font(name=FONT, size=11)
            c.alignment = Alignment(horizontal="center" if i != 4 else "right",
                                    vertical="center", wrap_text=True)
            c.border = BORDER
            c.fill = PatternFill("solid", fgColor=IDCOL if i == 2 else (BAND if r % 2 == 1 else "FFFFFF"))
        lc = ws2.cell(r, 5)
        lc.hyperlink = f"../Firmware/Modules/{mod}/{fname}"
        lc.font = Font(name=FONT, size=11, color="0563C1", underline="single")
        r += 1
    for i, w in enumerate([6, 12, 30, 44, 16], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ---------- راهنما ----------
    ws3 = wb.create_sheet("راهنما", 0)
    ws3.sheet_view.rightToLeft = True
    ws3.sheet_properties.tabColor = "7F7F7F"
    style_title(ws3, "فایل ولیدیشن کلی — راهنما", 2)
    lines = [
        ("این فایل چیست؟", "این فایل، نمای کلی همهٔ ولیدیشن‌هاست. از اینجا می‌بینی هر ماژول چند تست دارد و چندتایش انجام شده است."),
        ("تست‌ها کجاست؟", "تست‌ها اینجا تکرار نشده‌اند. هر تست فقط در فایل ولیدیشن ماژول خودش است. نتیجه را هم همان‌جا ثبت کن."),
        ("از کجا شروع کنم؟", "اول شیت «مرحله اول» را انجام بده. آن‌ها تست‌های راه‌اندازی اولیهٔ دستگاه‌اند (اولویت «بالا»). بعد بقیه را در فایل هر ماژول ادامه بده."),
        ("عددها تازه نیست؟", "عددهای «نمای کلی» لحظه‌ای‌اند. هر وقت وضعیت‌ها در فایل‌های ماژول عوض شد، این دستور را بزن تا این فایل تازه شود:"),
        ("دستور تازه‌سازی", "python3 tools/make_overall_validation.py"),
        ("رنگ‌ها", "همهٔ اکسل‌های پروژه یک رنگ‌بندی دارند: عنوان سرمه‌ای (1F3864)، سرستون آبی (1F4E79)، ستون شناسه آبی روشن (D9E2F3) و خانه‌های زرد (FFF2CC) یعنی «اینجا را تو پر کن / عوض کن». سبز یعنی سالم و قبول، زرد یعنی هشدار یا در انتظار، قرمز یعنی خطا یا رد."),
        ("تست‌های سیستمی", "چند تست سطح سیستم (روشن‌شدن کلی و صف احیا و پایه‌های ممنوعه) در فایل Documentation/System_State_Scenarios.xlsx است و در «نمای کلی» ردیف جدا دارد."),
    ]
    r = 3
    for k, v in lines:
        a = ws3.cell(r, 1, k)
        a.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        a.fill = PatternFill("solid", fgColor=HEAD)
        a.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        a.border = BORDER
        b = ws3.cell(r, 2, v)
        b.font = Font(name=FONT, size=11)
        b.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        b.border = BORDER
        ws3.row_dimensions[r].height = 42
        r += 1
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 95

    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print("wrote", OUT.relative_to(ROOT), "| stage1:", len(stage1))


if __name__ == "__main__":
    build()
