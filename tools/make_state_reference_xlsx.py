#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ساخت دو Workbook مرجع «حالت‌ها» و «سناریوها» از کد فعلی.

[EN] Regenerates Documentation/State.xlsx and
      Documentation/System_State_Scenarios.xlsx from the CURRENT code
      numbers (user order 2026-10-24: "yes, update the Excels with the
      current code"). Like every other workbook in this project these two
      are TOOL-GENERATED: edit this script and rerun it - never edit the
      xlsx files by hand.

      Documentation/System_State_Scenarios.xlsx is also an INPUT of
      tools/make_overall_validation.py (its Test_Cases sheet is counted
      there), so the sheet names/shape are kept stable.

Usage:
      python3 tools/make_state_reference_xlsx.py
"""

import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent

HEAD = "1F4E79"
TITLE = "1F3864"
IDCOL = "D9E2F3"
CELL = "FFF2CC"
BAND = "F2F2F2"
FONT = "B Nazanin"

BORDER = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def sheet_title(ws, text, cols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    c = ws.cell(1, 1, text)
    c.font = Font(name=FONT, size=13, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=TITLE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26


def sheet_header(ws, row, cols):
    for i in range(1, cols + 1):
        c = ws.cell(row, i)
        c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def put_row(ws, row, values, id_col=1, wrap=True):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row, i, v)
        c.font = Font(name=FONT, size=11)
        c.alignment = Alignment(horizontal="center" if i != id_col else "right",
                                vertical="center", wrap_text=wrap)
        c.border = BORDER
        c.fill = PatternFill("solid", fgColor=IDCOL if i == id_col else
                             (BAND if row % 2 == 1 else "FFFFFF"))


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ======================================================================
# Documentation/State.xlsx — ورودی/خروجی و جدول چراغ‌ها و بوق (از کد فعلی)
# ======================================================================
def build_state():
    wb = Workbook()

    ws = wb.active
    ws.title = "INPUT_OUTPUT"
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "C00000"
    sheet_title(ws, "ورودی‌ها و خروجی‌های دستگاه — بازسازی از کد فعلی (۲۰۲۶-۱۰-۲۴)", 4)

    ws.cell(2, 1, "INPUTS")
    ws.cell(2, 1).font = Font(name=FONT, size=11, bold=True)
    for i, h in enumerate(["ردیف", "ورودی", "واحد", "توضیحات"], start=1):
        ws.cell(3, i, h)
    sheet_header(ws, 3, 4)
    put_row(ws, 4, [1, "21..28", "V",
                    "برق ورودی مستقیم (DC). پنجرهٔ سالم ۲۱۰۰۰..۲۸۰۰۰ میلی‌ولت است؛ "
                    "بیرون از آن خطای ولتاژ/جریان ثبت می‌شود."])
    put_row(ws, 5, [2, "24 (دو باتری ۱۲ سری)", "V",
                    "دو باتری ۱۲ ولتی سری. هر نیم‌باتری جدا اندازه‌گیری می‌شود؛ "
                    "جمع دو نیم = ولتاژ پک."])

    ws.cell(7, 1, "OUTPUTS")
    ws.cell(7, 1).font = Font(name=FONT, size=11, bold=True)
    for i, h in enumerate(["ردیف", "خروجی", "واحد", "توضیحات"], start=1):
        ws.cell(8, i, h)
    sheet_header(ws, 8, 4)
    put_row(ws, 9, [1, "پک/ورودی", "V",
                    "خروجی سوییچ‌شده بین برق ورودی و باتری. تنها پایهٔ فرمان مسیر، "
                    "PB11 است (پایه‌های PB5/PB7 هرگز تغییر نمی‌کنند)."])
    put_row(ws, 10, [2, "دو کانال شارژ", "mA",
                     "شارژ هر نیم‌باتری با کانال خودش؛ هر کانال با کلید «فعال‌بودن شارژر» "
                     "قطع/وصل می‌شود."])
    widths(ws, [6, 22, 8, 70])

    ws2 = wb.create_sheet("UI")
    ws2.sheet_view.rightToLeft = True
    ws2.sheet_properties.tabColor = "C00000"
    sheet_title(ws2, "جدول چراغ‌ها و بوق — رفتار واقعی برنامه (Ui/README + کد فعلی)", 7)
    heads = ["ردیف", "حالت", "شرط", "سبز", "زرد", "قرمز", "بوق"]
    for i, h in enumerate(heads, start=1):
        ws2.cell(2, i, h)
    sheet_header(ws2, 2, 7)
    rows = [
        [1, "INPUT_OK / فول", "ورودی وصل و شارژ کامل",
         "ثابت", "خاموش", "خاموش", "ساکت"],
        [2, "CHARGING", "ورودی وصل و کانالی در شارژ است",
         "ثابت", "چشمک نسبت به مانده تا فول (هر چه پرتر، روشنی کوتاه‌تر، کف ۱۵۰ms)",
         "خاموش", "ساکت"],
        [3, "BatteryRun ≥ ۴۰٪", "برق شهر قطع، شارژ کافی",
         "چشمک با دورهٔ ۱۰۰۰ms (روشنی = درصد × ۱۰ms)", "خاموش", "خاموش", "ساکت"],
        [4, "BatteryRun ۲۰..۴۰٪", "برق شهر قطع",
         "همان چشمک", "خاموش", "خاموش", "یک بوق هر ۶۰s (هر بوق ۱۲۰۰ms)"],
        [5, "BatteryRun ۱۰..۲۰٪", "برق شهر قطع",
         "همان چشمک", "خاموش", "خاموش", "دو بوق هر ۶۰s (هر بوق ۱۱۵۰ms)"],
        [6, "BatteryRun ۱..۱۰٪", "برق شهر قطع",
         "چشمک سبز کوتاه", "خاموش", "خاموش", "سه بوق هر ۲۰s (هر بوق ۲۰۰۰ms)"],
        [7, "BatteryRun زیر ۱٪ (بحرانی)", "شارژ به آخرین حد رسیده",
         "خاموش پس از یک بوق", "خاموش", "خاموش",
         "یک بوق بلند ۱۰ثانیه‌ای فقط یک بار؛ بعد همهٔ چراغ‌ها خاموش تا خروج از ناحیهٔ بحرانی"],
        [8, "اضافه‌ولتاژ ورودی", "بیشتر از ۲۸۰۰۰ میلی‌ولت (پاک: ≤ ۲۷۰۰۰)",
         "ثابت", "خاموش", "پالس ۵۰٪ با دورهٔ ۱s", "هر ۱۰s یک بوق ۱s"],
        [9, "قطع باتری (BAT_LOST)", "پمپ > ۱۴٫۸V حین شارژ یا هر نیم < ۶V با ورودی سالم",
         "ثابت", "خاموش", "چشمک‌تند ۵۰/۵۰ دورهٔ ۱s", "سه بیپ کوتاه، دورهٔ تکرار ۳s"],
        [10, "سناریوی ۵ · عدم‌توازن (قفل‌شده)", "فرق دو نیم‌باتری بیش از حد",
         "ثابت", "خاموش", "چشمک‌زن با دوره و سهم مستقل",
         "دوره‌ای با دورهٔ مستقل (پیش‌فرض هر ۱ ساعت)"],
        [11, "سناریوی ۶ · باتری خراب", "۲۴ ساعت شارژ بدون یک‌بار فول",
         "ثابت", "خاموش", "ثابت", "بوق مستقل سناریو"],
        [12, "سناریوی ۷ · فنی", "خطای سخت‌افزاری مسیر قدرت",
         "چشمک همگام", "چشمک همگام", "چشمک همگام", "بوق مستقل سناریو"],
        [13, "سناریوی ۸ · باتری ضعیف", "فقط وقتی برق شهر قطع است",
         "خاموش", "خاموش", "چشمک‌زن (پیش‌فرض دوره ۵۰۰ms، سهم ۵۰٪)",
         "سه بوق هر ۳s (هر بوق ۲۰۰ms، گپ ۱۰۰ms)"],
        [14, "استراحت حرارتی", "۱ ساعت شارژ مداوم با جریان بالای آستانه",
         "ثابت", "چشمک‌تند", "خاموش", "ساکت؛ پس از ۱۰ دقیقه خودش ادامه می‌دهد"],
        [15, "بدون دادهٔ معتبر", "اندازه‌گیری یا قاب نامعتبر",
         "خاموش", "خاموش", "خاموش", "خاموش (حالت امن)"],
    ]
    for k, row in enumerate(rows):
        put_row(ws2, 3 + k, row, id_col=2)
    widths(ws2, [6, 26, 34, 26, 30, 22, 40])

    out = ROOT / "Documentation" / "State.xlsx"
    wb.save(out)
    print("wrote", out)


# ======================================================================
# Documentation/System_State_Scenarios.xlsx — مرجع سطح‌بالای State/Scenario
# ======================================================================
def build_scenarios():
    wb = Workbook()

    # ---------------- راهنما ----------------
    ws = wb.active
    ws.title = "راهنما"
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = "7F7F7F"
    sheet_title(ws, "راهنمای استفاده از مرجع State و Scenario", 3)
    for i, h in enumerate(["موضوع", "توضیح", ""], start=1):
        ws.cell(2, i, h)
    sheet_header(ws, 2, 3)
    guide = [
        ("راهنمای استفاده",
         "این Workbook مرجع سطح‌بالای State و سناریوهای سیستم ChangeOver است. "
         "جزئیات اجرایی هر ماژول در Workbookهای Validation همان ماژول ثبت می‌شود."),
        ("روش کار ۱",
         "ابتدا States و Transitions را بخوان؛ این دو برگه می‌گویند سیستم در هر شرط به کجا می‌رود."),
        ("روش کار ۲",
         "تست‌های ماژول‌ها در فایل ولیدیشن خودشان است (Firmware/Modules/<Module>) و "
         "اینجا تکرار نمی‌شود؛ در Test_Cases فقط تست‌های سطح سیستم (SYS/BZ/SAFE) هستند."),
        ("روش کار ۳",
         "در Coverage مطمئن شو هر الزام حداقل یک تست دارد؛ وضعیت GAP نباید بدون توضیح بماند."),
        ("اصل مهم",
         "snapshot نامعتبر قبل از هر تصمیم state، fault، timer یا pin بررسی می‌شود."),
        ("Build فعلی",
         "MODULE_UI=1، MODULE_MEASUREMENT=1، MODULE_CHANGEOVER=1، MODULE_FAULT=1، "
         "MODULE_CHARGER=1 (CHG_MASTER_ENABLE=1، CHG_TRANSFORMER_KNOWN=1)، "
         "MODULE_IMBALANCE=1، MODULE_ESPLINK=1، MODULE_JITTER=1، MODULE_PROTECTION=0، "
         "MODULE_MCU_POWER_PATH=1، MODULE_DIAG=1، MODULE_FLASH_LOCK=1."),
        ("محدودیت تست",
         "تست Host و syntax جای تست واقعی برد نیستند. برای Fault با Protection خاموش، "
         "fault باید از مسیر تست تزریق شود."),
        ("فایل‌های جزئیات",
         "UI_Board_Validation.xlsx، Measurement_Board_Validation.xlsx، "
         "Changeover_Board_Validation.xlsx و بقیهٔ Workbookهای ماژول‌ها."),
        ("تعریف PASS",
         "فقط وقتی PASS ثبت شود که نتیجهٔ واقعی با معیار قبولی ردیف منطبق باشد و مدرک تست ثبت شده باشد."),
        ("تازه‌سازی",
         "این فایل با کد تولید می‌شود؛ دستی ویرایش نکن. دستور: "
         "python3 tools/make_state_reference_xlsx.py"),
        ("تاریخ تهیه", "2026-09-16 | بازسازی از کد فعلی: 2026-10-24"),
    ]
    r = 3
    for k, v in guide:
        put_row(ws, r, [k, v], id_col=1)
        r += 1
    widths(ws, [22, 90, 2])

    # ---------------- States ----------------
    ws = wb.create_sheet("States")
    ws.sheet_view.rightToLeft = True
    sheet_title(ws, "States — حالت‌های رسمی سیستم (Changeover + UI)", 6)
    for i, h in enumerate(["State", "معنی", "ورودی معتبر لازم", "خروجی/رفتار مورد انتظار",
                           "شرط خروج", "تست مرجع"], start=1):
        ws.cell(2, i, h)
    sheet_header(ws, 2, 6)
    states = [
        ["BOOT", "شروع یا مقداردهی اولیه", "خیر/هنوز snapshot معتبر نیست",
         "state اولیه؛ PB11 در حالت Safe High", "snapshot معتبر و input_present یا fault", "—"],
        ["BATTERY", "ورودی اصلی حاضر نیست و مسیر باتری برقرار است",
         "بله، valid=true و input_present=false", "PB11 High؛ UI در BatteryRun",
         "cut condition یا input_present=true", "CO-006"],
        ["INPUT", "ورودی اصلی حاضر است", "بله، valid=true و input_present=true",
         "PB11 High؛ UI در InputOk/Charging", "cut condition یا input_present=false", "CO-007"],
        ["SAFE", "مسیر باتری قطع شده است", "بله برای ورود؛ در invalid حفظ می‌شود",
         "PB11 Low؛ تا reconnect قطع باقی می‌ماند",
         "input_present=true و v_bat24≥21200 (یا بازیابی بدون برق شهر — ST-011)", "CO-013..016"],
        ["FAULT", "fault معتبر در snapshot معتبر", "snapshot.valid=true و fault_mask!=0",
         "state=FAULT؛ هیچ pin توسط Fault تغییر نمی‌کند", "fault پاک شود", "CO-004, CO-005"],
    ]
    for k, row in enumerate(states):
        put_row(ws, 3 + k, row, id_col=1)
    widths(ws, [12, 34, 30, 38, 40, 16])

    # ---------------- Transitions ----------------
    ws = wb.create_sheet("Transitions")
    ws.sheet_view.rightToLeft = True
    sheet_title(ws, "Transitions — مسیرهای تغییر حالت (اعداد از کد فعلی)", 8)
    for i, h in enumerate(["ID", "From", "Trigger / guard", "Timer", "Action / output",
                           "To", "Validation", "توضیح"], start=1):
        ws.cell(2, i, h)
    sheet_header(ws, 2, 8)
    transitions = [
        ["ST-001", "هر state", "snapshot=NULL یا valid=false", "فوری",
         "cut/reconnect timer reset؛ state و PB11 حفظ", "همان state", "CO-001, CO-002, CO-017",
         "invalid قبل از هر تصمیمی"],
        ["ST-002", "BOOT", "valid=true و input_present=false", "فوری", "بدون cut؛ PB11 High",
         "BATTERY", "CO-006", "mapping معتبر"],
        ["ST-003", "BOOT", "valid=true و input_present=true", "فوری", "بدون cut؛ PB11 High",
         "INPUT", "CO-007", "mapping معتبر"],
        ["ST-004", "BATTERY/INPUT", "v_bat24<21000 و قفل باتری کم", "3000ms پیوسته",
         "BSP_GPIO_PROTECT_BATTERY=true", "SAFE", "CO-008", "cut گیت‌شده؛ ۲۱۰۰۰ دقیق"],
        ["ST-005", "BATTERY/INPUT", "v_bat24<20800، مستقل از قفل", "3000ms پیوسته",
         "BSP_GPIO_PROTECT_BATTERY=true", "SAFE", "CO-009", "cut مستقل؛ ۲۰۸۰۰ دقیق"],
        ["ST-006", "SAFE", "input_present=true و v_bat24≥21200", "30s پایداری ورودی + 3000ms ولتاژ",
         "BSP_GPIO_PROTECT_BATTERY=false", "INPUT", "CO-013, CO-014",
         "بازگشت با برق شهر؛ ۲۱۲۰۰ آستانه"],
        ["ST-007", "هر state معتبر", "fault_mask!=0", "فوری", "state=FAULT؛ pin بدون تغییر",
         "FAULT", "CO-004, CO-005", "fault اولویت دارد"],
        ["ST-008", "هر state", "snapshot.invalid و fault_mask!=0", "فوری",
         "fault نادیده؛ timer reset؛ pin حفظ", "همان state", "CO-003", "invalid قبل از fault"],
        ["ST-009", "هر state معتبر", "MillisecondsToTicks(3000)==0", "فوری",
         "هیچ cut/reconnect فوری؛ timer reset", "همان state", "CO-022", "guard صفر"],
        ["ST-010", "FAULT", "valid=true و fault_mask=0", "فوری",
         "mapping عادی یا SAFE بر اساس protect", "INPUT/BATTERY/SAFE", "CO-004 تکمیلی",
         "بسته به protect و input"],
        ["ST-011", "SAFE (بدون برق شهر)", "v_bat24≥21200 بی‌وقفه (۱۰s، هر افت = صفر) + فیلتر 3000ms",
         "10000ms + 3000ms", "مسیر باتری وصل می‌شود (یک تلاش)", "BATTERY", "CO-015/016",
         "بازیابی بدون برق شهر (۲۰۲۶-۱۰-۲۳)"],
        ["ST-012", "BATTERY (پس از ST-011)", "قطع دوباره داخل ۶۰s", "60000ms",
         "حکم باتری ضعیف (سناریوی ۸) قفل می‌شود؛ مسیر بسته", "SAFE + حکم ۸", "—",
         "محافظ چرخه؛ فقط برق شهر (۳۰s+3s) یا پاک‌کردن حکم باز می‌کند"],
        ["ST-013", "SAFE + حکم ۸", "v_bat24≥21200 پیوسته ۱۰ دقیقه (یا «ریست قفل‌ها»)",
         "600000ms", "حکم پاک و یک تلاش مجدد", "BATTERY", "—",
         "پاک‌شدن خودکار هشدار باتری فرسوده"],
    ]
    for k, row in enumerate(transitions):
        put_row(ws, 3 + k, row, id_col=1)
    widths(ws, [9, 18, 40, 22, 40, 18, 16, 34])

    # ---------------- Requirements ----------------
    ws = wb.create_sheet("Requirements")
    ws.sheet_view.rightToLeft = True
    sheet_title(ws, "Requirements — الزامات قابل ردیابی", 6)
    for i, h in enumerate(["Requirement", "Area", "Requirement دقیق", "منبع/ثابت",
                           "تست پوشش‌دهنده", "وضعیت"], start=1):
        ws.cell(2, i, h)
    sheet_header(ws, 2, 6)
    reqs = [
        ["REQ-M-001", "Measurement", "valid بعد از سه فریم کامل و پایدار true شود",
         "MEASUREMENT_WARMUP_FRAME_COUNT", "MS-001..006", "در انتظار برد"],
        ["REQ-M-002", "Measurement", "فریم ناقص warm-up را reset کند", "BSP ADC stable frame",
         "MS-005, MS-018", "در انتظار برد"],
        ["REQ-M-003", "Measurement", "snapshot اتمیک و GetSnapshot(NULL)=false", "measurement.c",
         "MS-015, MS-016", "در انتظار برد"],
        ["REQ-UI-001", "UI", "snapshot invalid/NULL باعث safe-off و flag=false شود", "ui_led.c",
         "UI-023..025, UI-035", "در انتظار برد"],
        ["REQ-UI-002", "UI", "Battery flag با 21000/21200 hysteresis و level پیوسته",
         "UI_LOW_BATTERY_..", "UI-027..030, UI-037", "در انتظار برد"],
        ["REQ-BZ-001", "Buzzer", "zero معتبر خاموش؛ ورودی نامعتبر -1 و Low", "ui_buzzer.h",
         "BZ-003..009", "در انتظار برد"],
        ["REQ-BZ-002", "Buzzer", "الگوی ON/gap/tail/nextCheck دقیق و non-blocking", "ui_buzzer.c",
         "BZ-001, BZ-010..022", "در انتظار برد"],
        ["REQ-CO-001", "Changeover", "invalid قبل از state/fault/pin و timerها reset شوند",
         "changeover.c", "CO-001..003, CO-017", "در انتظار برد"],
        ["REQ-CO-002", "Changeover", "fault معتبر → FAULT بدون تغییر pin", "fault_mask",
         "CO-004, CO-005, CO-018", "نیازمند injection"],
        ["REQ-CO-003", "Changeover", "mapping BATTERY/INPUT و SAFE/INPUT درست باشد",
         "changeover.c", "CO-006..016", "در انتظار برد"],
        ["REQ-CO-004", "Changeover", "cut گیت‌شده و مستقل هر دو 3000ms پیوسته",
         "21000/20800/3000ms", "CO-008..012", "در انتظار برد"],
        ["REQ-CO-005", "Changeover", "reconnect با input و ≥21200 بعد از 30s+3s",
         "21200/30000ms/3000ms", "CO-013..016", "در انتظار برد"],
        ["REQ-CO-006", "Safety", "فقط BSP_GPIO_PROTECT_BATTERY؛ PB5/PB7/HAL ممنوع",
         "changeover.c", "CO-019, CO-020", "بازبینی کد + برد"],
        ["REQ-CO-007", "Changeover",
         "بازیابی بدون برق شهر: ۱۰s سالمِ بی‌وقفه + 3s؛ قطع دوباره داخل ۶۰s = باتری ضعیف؛ "
         "پاک‌شدن حکم با ۱۰ دقیقه سالم یا ریست دستی",
         "147..157 (weak) + 150/149/148", "تست‌های سناریوی ۸", "در انتظار برد"],
        ["REQ-LOG-001", "Panel log",
         "تاریخچهٔ وضعیت همیشه می‌نویسد (حتی بدون برد)، ردیف قطع/وصل برق شهر حداکثر "
         "۲۵۰ms بعد می‌نشیند، هر خط وضعیت باتری (درصد/حالت/وات) را می‌گوید و در دشارژ "
         "هر بالا آمدنِ پک (پلهٔ ۲۰۰ میلی‌ولتی) ردیف خودش را می‌نویسد",
         "SIMPLE_LOG_MAINS_FLUSH_MS/ESP_LOG_MAINS_FLUSH_MS + ستون batt_pct + قانون VSTEP صعودی",
         "تست‌های هاست لاگ", "انجام شد (Host)"],
    ]
    for k, row in enumerate(reqs):
        put_row(ws, 3 + k, row, id_col=1)
    widths(ws, [14, 14, 56, 30, 24, 18])

    # ---------------- Test_Cases ----------------
    ws = wb.create_sheet("Test_Cases")
    ws.sheet_view.rightToLeft = True
    sheet_title(ws, "Test Cases — سناریوهای اجرایی سطح سیستم", 10)
    ws.cell(2, 1, "تست‌های هر ماژول فقط در فایل ولیدیشن همان ماژول است (پوشهٔ "
                  "Firmware/Modules) و اینجا تکرار نمی‌شود؛ این شیت فقط تست‌های سطح "
                  "سیستم را نگه می‌دارد.")
    ws.cell(2, 1).font = Font(name=FONT, size=10)
    for i, h in enumerate(["ID", "Area", "Setup / input", "روش اجرا", "نتیجهٔ مورد انتظار",
                           "معیار قبولی", "مرجع جزئیات", "Status", "نتیجهٔ واقعی", "مدرک"],
                          start=1):
        ws.cell(3, i, h)
    sheet_header(ws, 3, 10)
    cases = [
        ["SYS-001", "Startup", "Build فعال با UI/Measurement/Changeover/Fault=1",
         "Clean build، پروگرام، Reset", "بدون HardFault؛ taskهای فعال start شوند",
         "Reset و اجرای پیوسته", "—", "در انتظار", "", ""],
        ["SYS-002", "State log", "پنل با/بدون برد؛ قطع و وصل شارژر",
         "تغییر حالت‌ها و نگاه به تاریخچهٔ وضعیت",
         "ردیف بوت، ردیف قطع/وصل برق شهر (≤۲۵۰ms)، ضربان ۶ساعته",
         "هیچ تغییری بدون ردیف نماند", "تست‌های هاست لاگ", "انجام شد (Host)", "", ""],
        ["BZ-001", "Buzzer", "period=10000,duty=10,count=1,gap=0", "15s با scope/timer",
         "ON=1000 و tail=9000", "انحراف مجاز طبق Excel", "Buzzer BZ-001", "در انتظار", "", ""],
        ["BZ-003", "Buzzer", "period=10000,duty=0", "الگو را اجرا کن", "Low و return=0",
         "خاموشی امن", "Buzzer BZ-003", "در انتظار", "", ""],
        ["BZ-006", "Buzzer", "period=500,duty=50,count=1", "الگو را اجرا کن", "Low و return=-1",
         "نامعتبر رد شود", "Buzzer BZ-006", "در انتظار", "", ""],
        ["BZ-017", "Buzzer", "الگوی non-blocking فعال", "Measurement را همزمان مانیتور کن",
         "Measurement هر 10ms ادامه یابد", "task قفل نشود", "Buzzer BZ-017", "در انتظار", "", ""],
        ["SAFE-001", "GPIO", "برد در تمام سناریوها", "PB5/PB7/PB11 را مانیتور کن",
         "فقط PB11 در cut/reconnect تغییر کند", "PB5/PB7 بدون تغییر", "CO-019/020",
         "در انتظار", "", ""],
    ]
    for k, row in enumerate(cases):
        put_row(ws, 4 + k, row, id_col=1)
    widths(ws, [10, 12, 34, 30, 40, 26, 18, 16, 14, 12])

    # ---------------- Coverage ----------------
    ws = wb.create_sheet("Coverage")
    ws.sheet_view.rightToLeft = True
    sheet_title(ws, "Coverage — پوشش الزامات با تست‌ها", 5)
    for i, h in enumerate(["Requirement", "Covered by tests", "Expected evidence",
                           "Status", "Gap / action"], start=1):
        ws.cell(2, i, h)
    sheet_header(ws, 2, 5)
    cov = [
        ["Measurement warm-up", "MS-001..005", "Live snapshot + timing", "در انتظار",
         "سه فریم کامل و فریم ناپایدار"],
        ["Measurement conversion", "MS-004, MS-007, MS-020, MS-021", "DMM + raw counts + live values",
         "در انتظار", "مقادیر واقعی ثبت شود"],
        ["UI safe-off", "UI-023, UI-024", "LED/Buzzer voltage and flag", "در انتظار",
         "NULL/invalid تزریق شود"],
        ["UI hysteresis", "UI-027, UI-029", "flag trace", "در انتظار", "21000/21200 تست شود"],
        ["Buzzer timing", "BZ-001, BZ-010..017", "scope/timer CSV", "در انتظار",
         "مدت ON/gap/tail ثبت شود"],
        ["Changeover mapping", "CO-006, CO-007", "state + PB11", "در انتظار", "BATTERY/INPUT ثبت شود"],
        ["Changeover cut", "CO-008..012", "time + PB11 + state", "در انتظار", "2500/3000ms و مرزها"],
        ["Changeover reconnect", "CO-013..016", "time + PB11 + state", "در انتظار",
         "21200 و کمتر از 3s"],
        ["Weak-battery recovery", "سناریوی ۸", "time + face + one retry", "در انتظار",
         "10s hold، 60s cycle guard، 10min clear"],
        ["State log completeness", "SYS-002 + تست‌های هاست", "فایل‌های روزانه + CSV",
         "انجام شد (Host)", "روی برد واقعی تکرار شود"],
        ["Fault valid", "FL-001", "fault injection + state/pin", "نیازمند injection",
         "Protection خاموش است"],
        ["Fault invalid", "FL-002", "fault injection + preserved state/pin", "نیازمند injection",
         "invalid باید قبل از fault باشد"],
        ["GPIO restriction", "SAFE-001", "scope/logic analyzer", "در انتظار", "PB5/PB7 بدون تغییر"],
    ]
    for k, row in enumerate(cov):
        put_row(ws, 3 + k, row, id_col=1)
    widths(ws, [26, 28, 34, 18, 34])

    # ---------------- Signals ----------------
    ws = wb.create_sheet("Signals")
    ws.sheet_view.rightToLeft = True
    sheet_title(ws, "Signals — سیگنال‌ها و مالکیت داده", 6)
    for i, h in enumerate(["Signal / variable", "Owner", "Meaning", "Readers",
                           "Writer / source", "Safe/default"], start=1):
        ws.cell(2, i, h)
    sheet_header(ws, 2, 6)
    sigs = [
        ["measurement_snapshot_t", "Measurement", "دادهٔ اتمیک ADC و input",
         "UI, Changeover, Protection, Charger, McuPowerPath", "Measurement task", "valid=false"],
        ["snapshot.valid", "Measurement", "آماده‌بودن داده بعد از warm-up",
         "UI, Changeover, Charger, Fault", "Measurement", "false"],
        ["snapshot.v_bat24_mv", "Measurement", "ولتاژ باتری 24V", "UI, Changeover", "Measurement", "0"],
        ["snapshot.input_present", "Measurement", "حضور ورودی اصلی", "UI, Changeover",
         "BSP/Measurement", "false"],
        ["BOOL__G__UiBatteryAlarmIssued", "UI", "آلارم باتری کم با hysteresis", "Changeover",
         "UI", "false"],
        ["fault_mask", "Fault", "خطاهای مشترک لچ‌شده", "Changeover, Charger, UI, Comm",
         "Fault producers", "FAULT_NONE"],
        ["BSP_GPIO_PROTECT_BATTERY", "BSP/Changeover", "PB11 منطقی حفاظت باتری", "BSP",
         "Changeover only", "false = Safe High"],
        ["PB5/PB7", "BSP", "پایه‌های خارج از مالکیت Changeover", "هیچ‌کدام در Changeover",
         "ماژول‌های مجاز خودشان", "بدون تغییر"],
        ["CH_PD (ESP power)", "EspLink", "تغذیهٔ ماژول وای‌فای با سیاست صرفه‌جویی",
         "—", "EspLink_Power (با مهلت ۴۰۰ms برای آخرین فریم)", "ON"],
        ["تاریخچهٔ وضعیت (پنل)", "Panel (ESP)", "ردیف‌های تغییر وضعیت روی LittleFS",
         "کارت تاریخچه + دانلود CSV", "OnFrame/Tick پنل", "روز بدون ساعت"],
    ]
    for k, row in enumerate(sigs):
        put_row(ws, 3 + k, row, id_col=1)
    widths(ws, [30, 18, 40, 40, 36, 18])

    # ---------------- Faults ----------------
    ws = wb.create_sheet("Faults")
    ws.sheet_view.rightToLeft = True
    sheet_title(ws, "Faults — ماسک خطا و روش تزریق", 6)
    for i, h in enumerate(["Fault bit", "Intended producer", "Meaning", "Consumer",
                           "Test method now", "Recovery"], start=1):
        ws.cell(2, i, h)
    sheet_header(ws, 2, 6)
    faults = [
        ["FAULT_ADC", "Protection/Measurement", "ADC/snapshot fault", "Changeover",
         "func__Fault_Set injection؛ Protection فعلاً off", "Fault_Clear/test reset"],
        ["FAULT_OVERCURRENT_1", "Protection", "جریان کانال 1 بیش از حد", "Changeover",
         "injection؛ threshold کامل نشده", "Fault_Clear"],
        ["FAULT_OVERCURRENT_2", "Protection", "جریان کانال 2 بیش از حد", "Changeover",
         "injection؛ threshold کامل نشده", "Fault_Clear"],
        ["FAULT_LOW_BATTERY", "Protection/UI policy", "خطای کم‌باتری ایمنی", "Changeover",
         "injection؛ UI flag جداست", "Fault_Clear"],
        ["FAULT_JITTER_1", "Jitter", "رویداد jitter خط 1", "Charger (JIT retry), Comm",
         "روی برد: پالس LM393 → EXTI؛ در host: injection", "auto/retry"],
        ["FAULT_JITTER_2", "Jitter", "رویداد jitter خط 2", "Charger (JIT retry), Comm",
         "روی برد: پالس LM393 → EXTI؛ در host: injection", "auto/retry"],
        ["FAULT_CHARGER_BAT_LOST", "Fault",
         "قطع باتری: پمپ >14.8V حین شارژ (150ms) یا هر نیم <6V با ورودی سالم (1s)",
         "Charger, UI", "تزریق/سیم‌کشی", "باتری برگردد + 1s پایدار"],
        ["FAULT_CHARGER_TECHNICAL", "Charger (scenario 7)",
         "مسیر قدرت: duty>200‰ با جریان صفر (و بالعکس)", "UI, Panel", "تزریق",
         "خاموش/روشن تا ریست؛ قفل نهایی فقط با ریست برد"],
    ]
    for k, row in enumerate(faults):
        put_row(ws, 3 + k, row, id_col=1)
    widths(ws, [26, 22, 44, 22, 36, 30])

    # ---------------- Change_Log ----------------
    ws = wb.create_sheet("Change_Log")
    ws.sheet_view.rightToLeft = True
    sheet_title(ws, "Change Log — تاریخچهٔ تغییرات", 5)
    for i, h in enumerate(["Date", "Change", "Reason", "Source commit", "Validation"],
                          start=1):
        ws.cell(2, i, h)
    sheet_header(ws, 2, 5)
    log = [
        ["2026-09-16", "ساخت مرجع Master State و Scenario",
         "قابل‌فهم‌کردن stateها، transitionها و پوشش تست", "b64d7ed base",
         "این Workbook + draw.io"],
        ["2026-09-16", "وضعیت فعال‌سازی",
         "UI/Measurement/Changeover/Fault برای build تستی فعال؛ Protection و ماژول‌های آینده خاموش",
         "b64d7ed", "فلگ‌ها در modules_config"],
        ["2026-09-22", "ممیزی کامل برنامه (فیکس race در BspExti_TakeEvent، FAULT_ADC لحظه‌ای، "
         "حذف هوک تکراری ISR)؛ فیلترهای جریان median-3 → میانگین-10",
         "هماهنگی با کد", "—", "تست‌های هاست"],
        ["2026-10-24", "بازسازی هر دو Workbook (State.xlsx و همین فایل) از کد فعلی با ابزار "
         "tools/make_state_reference_xlsx.py: جدول چراغ‌ها/بوق، پنجرهٔ ورودی ۲۱..۲۸، "
         "قواعد وصل مجدد بدون برق شهر (ST-011..013) و الزام کامل‌بودن لاگ (REQ-LOG-001)",
         "دستور کاربر: «اکسل‌ها را با کد فعلی به‌روز کن»", "v1.108",
         "تست‌های هاست + تولید ابزار"],
    ]
    for k, row in enumerate(log):
        put_row(ws, 3 + k, row, id_col=1)
    widths(ws, [14, 60, 40, 18, 24])

    out = ROOT / "Documentation" / "System_State_Scenarios.xlsx"
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    build_state()
    build_scenarios()
    print("state reference workbooks written")
