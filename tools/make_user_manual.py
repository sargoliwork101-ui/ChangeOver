#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_user_manual.py - ساخت راهنمای کاربر (Word) برای برد ChangeOver.

[EN] Builds Documentation/User_Manual.docx: a simple, plain-language user
     manual (setup, panel, indicators, errors). No calibration, no internals.
     Image placeholders are written in RED so the owner can find, replace
     and delete them.
[FA] راهنمای کاربر ساده با جمله‌های کوتاه می‌سازد؛ جای عکس‌ها قرمز است.
Run: python3 tools/make_user_manual.py
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import pathlib

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / 'Documentation' / 'User_Manual.docx'

FONT = 'Vazirmatn'
NAVY = '1F3864'
RED = 'FF0000'


def _rtl_run(r, size=12, bold=False, color=None):
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), FONT)
    rf.set(qn('w:hAnsi'), FONT)
    rf.set(qn('w:cs'), FONT)
    rPr.append(rf)
    rtl = OxmlElement('w:rtl')
    rtl.set(qn('w:val'), '1')
    rPr.append(rtl)
    return r


def para(doc_or_cell, text='', size=12, bold=False, color=None, align='right',
         space_after=6, line=1.35):
    p = doc_or_cell.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)
    p.alignment = {'right': WD_ALIGN_PARAGRAPH.RIGHT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line
    if text:
        _rtl_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def heading(doc, text, level=1):
    p = para(doc, '', space_after=8 if level == 1 else 6)
    _rtl_run(p.add_run(text), size=17 if level == 1 else 14, bold=True, color=NAVY)
    return p


def item(doc, text, size=12):
    p = para(doc, '', space_after=3)
    _rtl_run(p.add_run('•  '), size=size, bold=True, color=NAVY)
    _rtl_run(p.add_run(text), size=size)
    return p


def step(doc, n, text, size=12):
    p = para(doc, '', space_after=3)
    _rtl_run(p.add_run(f'قدم {n}:  '), size=size, bold=True, color=NAVY)
    _rtl_run(p.add_run(text), size=size)
    return p


def photo(doc, text):
    p = para(doc, '', align='center', space_after=2)
    _rtl_run(p.add_run('【 جای عکس: ' + text + ' 】'), size=11, bold=True, color=RED)
    para(doc, '(این نوشتهٔ قرمز فقط راهنماست؛ عکس را اینجا بگذارید و بعد این متن را پاک کنید.)',
         size=10, bold=True, color=RED, align='center', space_after=12)
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t._tbl.tblPr.append(OxmlElement('w:bidiVisual'))

    def fill(cell, text, bold=False, color=None, size=11):
        cell.text = ''
        p = cell.paragraphs[0]
        pPr = p._p.get_or_add_pPr()
        b = OxmlElement('w:bidi')
        b.set(qn('w:val'), '1')
        pPr.append(b)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.line_spacing = 1.25
        _rtl_run(p.add_run(text), size=size, bold=bold, color=color)

    for j, h in enumerate(headers):
        fill(t.rows[0].cells[j], h, bold=True, color=NAVY)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            fill(cells[j], val)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    para(doc, '', space_after=4)
    return t


def build():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.2)
    doc.styles['Normal'].font.name = FONT
    doc.styles['Normal'].font.size = Pt(12)

    # ---------- cover ----------
    para(doc, '', space_after=30)
    para(doc, 'راهنمای کاربر', size=26, bold=True, color=NAVY, align='center', space_after=10)
    para(doc, 'برد شارژر و تعویض خودکار باتری (ChangeOver)', size=15, bold=True, align='center', space_after=6)
    para(doc, 'نسخهٔ ۱', size=12, align='center', space_after=24)
    photo(doc, 'عکس برد و دستگاه کامل')

    # ---------- 1 ----------
    heading(doc, '۱) این دستگاه چه کار می‌کند؟')
    para(doc, 'این دستگاه دو باتری ۱۲ ولتی را شارژ می‌کند و برق خروجی را بین برق ورودی و باتری جابه‌جا می‌کند.', align='justify')
    para(doc, 'وقتی برق ورودی هست، بار از برق ورودی کار می‌کند و باتری‌ها هم شارژ می‌شوند.', align='justify')
    para(doc, 'وقتی برق ورودی قطع شود، دستگاه خودش و بدون وقفه بار را روی باتری می‌برد.', align='justify')
    para(doc, 'وقتی برق ورودی برگردد، دستگاه خودش به برق ورودی برمی‌گردد و شارژ باتری‌ها ادامه پیدا می‌کند.', align='justify')
    para(doc, 'شما لازم نیست کاری بکنید؛ همه‌چیز خودکار است. چراغ‌ها به شما می‌گویند دستگاه چه وضعی دارد.', align='justify')

    # ---------- 2 ----------
    heading(doc, '۲) برای راه‌اندازی چه لازم است؟')
    item(doc, 'دو باتری ۱۲ ولتی سرب‌اسیدی که سری وصل می‌شوند (روی هم ۲۴ ولت).')
    item(doc, 'برق ورودی مستقیم (DC) بین ۲۲ تا ۲۸ ولت.')
    item(doc, 'بار یا مصرف‌کننده‌ای که باید همیشه روشن بماند.')
    item(doc, 'برای دیدن پنل: یک گوشی یا کامپیوتر با وای‌فای و مرورگر (مثل کروم).')

    # ---------- 3 ----------
    heading(doc, '۳) وصل‌کردن دستگاه')
    photo(doc, 'عکس برد با نشان‌دادن جای سیم باتری، ورودی و خروجی')
    step(doc, 1, 'اول دستگاه را خاموش کنید و برق ورودی را قطع کنید.')
    step(doc, 2, 'دو باتری را سری وصل کنید: مثبت باتری اول به منفی باتری دوم.')
    step(doc, 3, 'سیم‌های باتری را به جای باتری روی برد وصل کنید. مثبت و منفی را جابه‌جا نزنید.')
    step(doc, 4, 'برق ورودی را به جای ورودی روی برد وصل کنید.')
    step(doc, 5, 'بار را به خروجی وصل کنید.')
    step(doc, 6, 'یک بار دیگر همهٔ سیم‌ها را نگاه کنید که محکم و سر جای درستشان باشند.')

    # ---------- 4 ----------
    heading(doc, '۴) روشن‌کردن')
    step(doc, 1, 'برق ورودی را وصل و دستگاه را روشن کنید.')
    step(doc, 2, 'اولِ کار، هر سه چراغ برای چند لحظه روشن می‌شوند. این آزمایش خودِ دستگاه است و طبیعی است.')
    step(doc, 3, 'چند ثانیه صبر کنید؛ شارژ باتری‌ها خودش شروع می‌شود.')
    step(doc, 4, 'وقتی چراغ سبز ثابت شد، یعنی باتری پر است و دستگاه آماده است.')
    para(doc, 'از این به بعد دستگاه خودش کار می‌کند و لازم نیست هر روز به آن سر بزنید.', align='justify')

    # ---------- 5 ----------
    heading(doc, '۵) چراغ‌ها چه می‌گویند؟')
    para(doc, 'روی برد سه چراغ هست: سبز، زرد و قرمز. این جدول معنی هر حالت را می‌گوید:', align='justify')
    table(doc,
          ['چراغ', 'معنی', 'چه کار کنید؟'],
          [
              ['سبز ثابت', 'همه‌چیز خوب است؛ برق ورودی هست و باتری پر است.', 'کاری لازم نیست.'],
              ['زرد چشمک‌زن + سبز ثابت', 'باتری در حال شارژ است.', 'کاری لازم نیست؛ صبر کنید تا سبز ثابت شود.'],
              ['زرد چشمک‌زن تند + سبز ثابت', 'دستگاه دارد استراحت می‌کند تا داغ نشود.', 'طبیعی است؛ بعد از چند دقیقه خودش دوباره شارژ را شروع می‌کند.'],
              ['سبز چشمک‌زن (برق ورودی قطع است)', 'بار دارد با باتری کار می‌کند. تند و کند شدن چشمک، پر و خالی بودن باتری را نشان می‌دهد.', 'اگر می‌توانید، برق ورودی را زودتر برگردانید.'],
              ['قرمز چشمک‌زن + بوق', 'یک خطا پیش آمده است.', 'بخش ۸ «خطاها» را ببینید.'],
              ['قرمز ثابت', 'باتری خراب شده است.', 'باتری باید عوض شود.'],
              ['هر سه چراغ با هم چشمک‌زن + بوق', 'خطای سخت‌افزاری در خود دستگاه است.', 'دستگاه را یک بار خاموش و روشن کنید. اگر درست نشد، با پشتیبانی تماس بگیرید.'],
          ],
          widths=[4.4, 6.2, 5.4])
    photo(doc, 'عکس چراغ‌های روی برد در حالت عادی (سبز ثابت)')

    # ---------- 6 ----------
    heading(doc, '۶) صدای بوق چه می‌گوید؟')
    item(doc, 'در کار عادی، دستگاه ساکت است. بوق فقط برای خبرهای مهم است.')
    item(doc, 'وقتی دستگاه با باتری کار می‌کند و شارژ باتری خیلی کم شود، بوق هشدار می‌دهد؛ یعنی بهتر است برق ورودی زود برگردد.')
    item(doc, 'وقتی خطایی باشد، بوق با فاصله تکرار می‌شود تا متوجه شوید و چراغ‌ها را نگاه کنید.')

    # ---------- 7 ----------
    heading(doc, '۷) وقتی برق ورودی قطع می‌شود')
    para(doc, 'نگران نباشید؛ دستگاه همان لحظه بار را روی باتری می‌برد و کارتان قطع نمی‌شود.', align='justify')
    para(doc, 'چراغ سبز چشمک می‌زند تا بدانید دستگاه دارد با باتری کار می‌کند.', align='justify')
    para(doc, 'اگر شارژ باتری خیلی کم شود، اول بوق هشدار می‌دهد. اگر برق باز هم برنگردد، دستگاه خروجی را قطع می‌کند تا باتری خراب نشود.', align='justify')
    para(doc, 'با برگشتن برق ورودی، دستگاه خودش خروجی را وصل می‌کند و شارژ باتری‌ها را از سر می‌گیرد.', align='justify')

    # ---------- 8 errors ----------
    heading(doc, '۸) خطاها و راه‌حل‌ها')
    para(doc, 'اگر چراغ قرمز دیدید یا بوق شنیدید، این جدول را نگاه کنید:', align='justify')
    table(doc,
          ['نشانه', 'یعنی چه؟', 'چه کار کنید؟'],
          [
              ['قرمز چشمک‌زن یک‌ثانیه‌ای و سه بوق کوتاه، پشت سر هم',
               'سیم باتری قطع شده یا شل است.',
               'دستگاه را خاموش کنید. سیم‌های باتری و سر باتری‌ها را نگاه کنید؛ اگر شل یا کثیف‌اند درستشان کنید. بعد دستگاه را روشن کنید.'],
              ['قرمز ثابت',
               'باتری خراب است و دیگر خوب شارژ نمی‌شود.',
               'باتری را با باتری سالم عوض کنید.'],
              ['قرمز چشمک‌زن و هر چند وقت یک بوق',
               'دو نیمهٔ باتری با هم برابر نیستند؛ یکی ضعیف‌تر شده است.',
               'باتری نیاز به بررسی دارد. اگر این حالت ماند، باتری را عوض کنید یا با پشتیبانی تماس بگیرید.'],
              ['چهرهٔ اضافه‌ولتاژ (در پنل هم دیده می‌شود)',
               'ولتاژ برق ورودی خیلی بالا رفته است (بیشتر از ۲۸ ولت).',
               'منبع برق ورودی را بررسی کنید و ولتاژ درست (۲۲ تا ۲۸ ولت) بدهید.'],
              ['هر سه چراغ با هم چشمک‌زن و بوق',
               'داخل خود دستگاه ایراد سخت‌افزاری پیدا شده است.',
               'دستگاه را خاموش کنید، چند ثانیه صبر کنید و روشن کنید. اگر دوباره همان شد، دستگاه را خاموش نگه دارید و با پشتیبانی تماس بگیرید.'],
              ['زرد چشمک‌زن تند و سبز ثابت، مدت طولانی',
               'دستگاه برای خنک‌شدن استراحت می‌کند.',
               'طبیعی است و خطا نیست. صبر کنید؛ خودش برمی‌گردد.'],
          ],
          widths=[4.4, 5.2, 6.4])
    para(doc, 'نکته: بیشتر خطاها با درست‌شدن علت، خودشان پاک می‌شوند. مثلاً اگر سیم باتری را درست وصل کنید، خطای قطع باتری بعد از چند ثانیه می‌رود و شارژ از سر گرفته می‌شود.', align='justify')

    # ---------- 9 panel ----------
    heading(doc, '۹) پنل دستگاه')
    para(doc, 'دستگاه یک پنل بی‌سیم دارد. با گوشی یا کامپیوتر می‌توانید وضعیت را ببینید و چند کار ساده انجام دهید.', align='justify')
    heading(doc, 'وصل‌شدن به پنل', level=2)
    step(doc, 1, 'وای‌فای گوشی یا کامپیوتر را روشن کنید.')
    step(doc, 2, 'به شبکهٔ وای‌فای پنل وصل شوید. نام شبکه و رمز آن روی برچسب دستگاه نوشته شده است.')
    step(doc, 3, 'مرورگر را باز کنید و نشانی پنل را که روی همان برچسب است بنویسید.')
    step(doc, 4, 'صفحهٔ پنل باز می‌شود. اگر صفحه نیامد، یک بار مرورگر را ببندید و دوباره باز کنید.')
    photo(doc, 'عکس برچسب دستگاه که نام وای‌فای، رمز و نشانی پنل روی آن است')
    heading(doc, 'صفحهٔ وضعیت', level=2)
    item(doc, 'ولتاژ باتری‌ها و برق ورودی را زنده نشان می‌دهد.')
    item(doc, 'می‌گوید هر باتری در چه مرحله‌ای شارژ می‌شود: شارژ سریع، شارژ کامل‌کننده یا پر و آماده.')
    item(doc, 'اگر خطایی باشد، همین‌جا نوشته می‌شود.')
    item(doc, 'دو دکمهٔ روشن و خاموش برای دو شارژر دارد. با خاموش‌کردن، شارژ آن باتری همان لحظه قطع می‌شود. برای کار عادی، هر دو را روشن نگه دارید.')
    photo(doc, 'عکس صفحهٔ وضعیت پنل در گوشی')
    heading(doc, 'تنظیمات شارژ', level=2)
    item(doc, 'این صفحه برای نصاب و آدم فنی است. اگر لازم ندارید، چیزی را در آن عوض نکنید.')
    item(doc, 'هر عددی که عوض کنید، خودش ذخیره می‌شود و با قطع برق از بین نمی‌رود.')
    item(doc, 'اگر چیزی را اشتباه عوض کردید، دکمهٔ «بازگردانی پیش‌فرض کارخانه» همه را به حالت اول برمی‌گرداند.')
    item(doc, 'در همین پنل می‌توانید چراغ‌های دستگاه را هم زنده ببینید.')

    # ---------- 10 troubleshooting ----------
    heading(doc, '۱۰) اگر چیزی درست کار نکرد')
    table(doc,
          ['مشکل', 'اول این‌ها را نگاه کنید'],
          [
              ['هیچ چراغی روشن نمی‌شود', 'برق ورودی وصل است؟ سیم‌ها محکم‌اند؟ باتری‌ها وصل‌اند؟'],
              ['باتری شارژ نمی‌شود', 'برق ورودی بین ۲۲ تا ۲۸ ولت هست؟ در پنل، شارژر خاموش نشده باشد؟ سیم باتری شل نباشد؟'],
              ['با قطع برق ورودی، خروجی هم قطع شد', 'احتمالاً شارژ باتری خیلی کم شده و دستگاه برای محافظت از باتری قطع کرده است. برق ورودی را برگردانید و کمی صبر کنید.'],
              ['پنل باز نمی‌شود', 'به وای‌فای پنل وصل‌اید؟ نشانی را درست نوشته‌اید؟ یک بار وای‌فای را خاموش و روشن کنید.'],
              ['بوق می‌زند ولی معلوم نیست چرا', 'چراغ‌ها را نگاه کنید و در جدول بخش ۸ پیدایش کنید. در پنل هم خطا نوشته می‌شود.'],
          ],
          widths=[5.2, 10.8])

    # ---------- 11 safety ----------
    heading(doc, '۱۱) نکات ایمنی کوتاه')
    item(doc, 'مثبت و منفی باتری و ورودی را جابه‌جا نزنید.')
    item(doc, 'دو سر باتری را به هم نزنید (اتصال کوتاه).')
    item(doc, 'دستگاه را خیس نکنید و کنار گرما و آتش نگذارید.')
    item(doc, 'در دستگاه را باز نکنید؛ داخل آن چیزی برای تعمیر شما نیست.')
    item(doc, 'اگر بوی سوختگی یا دود دیدید، فوری برق ورودی و باتری را جدا کنید و با پشتیبانی تماس بگیرید.')

    # ---------- 12 support ----------
    heading(doc, '۱۲) پشتیبانی')
    para(doc, 'اگر با این راهنما مشکل حل نشد، با پشتیبانی تماس بگیرید و بگویید چراغ‌ها چه حالتی دارند؛ همین به فهمیدن مشکل کمک زیادی می‌کند.', align='justify')
    para(doc, '【 شماره و نشانی پشتیبانی را اینجا بنویسید 】', size=11, bold=True, color=RED, align='center')

    doc.save(str(OUT))
    print('wrote', OUT)


if __name__ == '__main__':
    build()
